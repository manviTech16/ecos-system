#!/usr/bin/env python3
"""Import the ECOS Excel workbook into a SQLite database.

Usage: python import_excel.py [workbook.xlsx] [database.db]
The database is rebuilt from scratch every run."""
import os
import sqlite3
import sys

from openpyxl import load_workbook

from schema import TABLES

HERE = os.path.dirname(os.path.abspath(__file__))
SQLT = {"text": "TEXT", "int": "INTEGER", "real": "REAL"}


def ddl(spec):
    lines = []
    for col, typ in spec["cols"].items():
        pk = " PRIMARY KEY" if col == spec["pk"] else " NOT NULL" if col in spec["required"] else ""
        lines.append(f'  "{col}" {SQLT[typ]}{pk}')
    for col, (route, pk) in spec["fks"].items():
        lines.append(f'  FOREIGN KEY("{col}") REFERENCES "{TABLES[route]["table"]}"("{pk}")')
    return f'CREATE TABLE "{spec["table"]}" (\n' + ",\n".join(lines) + "\n)"


def convert(value, typ):
    if value is None:
        return None
    if typ == "int":
        return int(float(value))
    if typ == "real":
        return float(value)
    return str(value).strip()


def run(xlsx, db):
    if not os.path.exists(xlsx):
        sys.exit(f"Workbook not found: {xlsx}")
    if os.path.exists(db):
        os.remove(db)
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    con = sqlite3.connect(db)
    con.execute("PRAGMA journal_mode=WAL")
    order = ["departments", "employees", "projects", "tasks", "assignments", "constraints", "communications", "meetings", "resources", "dependencies"]
    for route in order:
        spec = TABLES[route]
        con.execute(ddl(spec))
        if spec["table"] not in wb.sheetnames:
            sys.exit(f'Sheet "{spec["table"]}" is missing from the workbook')
        rows = list(wb[spec["table"]].iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        missing = [c for c in spec["cols"] if c not in header]
        if missing:
            sys.exit(f'Sheet "{spec["table"]}" is missing columns: {missing}')
        idx = [header.index(c) for c in spec["cols"]]
        data = [[convert(r[i], t) for i, t in zip(idx, spec["cols"].values())] for r in rows[1:] if r[0] is not None]
        con.executemany(f'INSERT INTO "{spec["table"]}" VALUES ({",".join("?" * len(idx))})', data)
        for col in spec["fks"]:
            con.execute(f'CREATE INDEX "ix_{spec["table"]}_{col}" ON "{spec["table"]}"("{col}")')
        print(f'  {spec["table"]:<18}{len(data):>6} rows')
    bad = con.execute("PRAGMA foreign_key_check").fetchall()
    con.commit()
    con.close()
    if bad:
        print(f"WARNING: {len(bad)} foreign-key problems found, e.g. {bad[:3]}")
    else:
        print("Foreign keys: OK")
    print(f"Database ready: {db}")


if __name__ == "__main__":
    run(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "data", "Enterprise_Collaboration_Optimization_Dataset.xlsx"),
        sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "ecos.db"))
