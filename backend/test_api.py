#!/usr/bin/env python3
"""End-to-end tests. Builds a throwaway database from the workbook, starts the server on a free port and calls the real API.
Run: python test_api.py"""
import importlib.util
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "data", "Enterprise_Collaboration_Optimization_Dataset.xlsx")
PROC = None
BASE = ""
BASELINE = {}
BASE_VIOL = -1


SAMPLE = 'Q3 2026 Workforce Productivity Report\n\nExecutive Summary\nThis report reviews workforce productivity and project delivery across all departments for the third quarter of 2026. It draws on task records, assignment data and resource logs for 150 employees working on 50 projects. Overall delivery improved, but workload is unevenly distributed and several high-priority tasks remain at risk.\n\nIntroduction\nManagement commissioned this review after a rise in delayed projects during the second quarter. The scope covers Software Development, QA, Data Analytics and Operations. Data was collected between July 1 and September 30, 2026.\n\nKey Findings\nTask completion rose to 62% in Q3, up from 54% in Q2, a gain of 8 percentage points. The Software Development department completed 1,240 tasks, which is 31% more than in the previous quarter. Average employee workload was 23%, but 12 employees carried more than 40% of their capacity. Our analysis revealed that 18% of blocked tasks were waiting on resources rather than people. GPU resources averaged 94% utilization, which exceeded the planning threshold of 85%. Customer-facing projects missed their deadlines 2.4 times more often than internal projects (p < 0.05). Contact the analytics lead at maria.silva@example.com or 415-555-0199 for the raw data.\n\nAnalysis\nThe data shows that delays concentrate in projects that depend on shared hardware. Teams with fewer than 5 members delivered 14% faster than larger teams. Reassigning tasks to employees with matching skills reduced average task duration from 9.1 days to 7.4 days.\n\nConclusions\nIn conclusion, productivity gains in Q3 were real but fragile. Resource contention, not headcount, is the main driver of delay. Therefore, the organisation should treat shared hardware as the critical constraint when planning Q4 delivery.\n\nRecommendations\nWe recommend adding two GPU nodes before the end of October to bring utilization below 85%.\nManagers should rebalance work away from the 12 overloaded employees within the next 30 days.\nThe PMO must review all blocked high-priority tasks every week.\nConsider adopting skill-based assignment for all new tasks, since it cut task duration by 19%.\n'
STUB_URL, STUB_LOG = "", []


def start_stub():
    """A tiny fake of the Anthropic Messages API so the AI path can be tested without a real key or network."""
    global STUB_URL
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    class Stub(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            doc = body["messages"][0]["content"].split("<document>")[1].split("</document>")[0]
            STUB_LOG.append(dict(headers=dict(self.headers), body=body, doc=doc))
            if "TRIGGER_FAIL" in doc:
                return self._send(500, {"error": "boom"})
            if "TRIGGER_AUTH" in doc:
                return self._send(401, {"error": "bad key"})
            text = "Sorry, I cannot help." if "TRIGGER_BADJSON" in doc else "```json\n" + json.dumps(dict(
                overview="This report reviews Q3 workforce productivity. Overall delivery improved.",
                key_points=["Task completion rose to 62% in Q3."], findings=["GPU resources averaged 94% utilization."],
                statistics=["Task completion rose to 62% in Q3.", "Revenue grew by 999% last year."],
                conclusions=["Resource contention is the main driver of delay."], recommendations=["Add two GPU nodes before October ends."])) + "\n```"
            self._send(200, {"content": [{"type": "text", "text": text}]})

        def _send(self, status, obj):
            data = json.dumps(obj).encode()
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Stub)
    STUB_URL = f"http://127.0.0.1:{srv.server_address[1]}"
    import threading
    threading.Thread(target=srv.serve_forever, daemon=True).start()


def upload(name, data, mode="local", extra_headers=None):
    req = urllib.request.Request(f"{BASE}/api/reports/summarize?mode={mode}", data=data, method="POST",
                                 headers={"Content-Type": "application/octet-stream", "X-File-Name": urllib.parse.quote(name), **(extra_headers or {})})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def make_docx(paras, extra_files=None, document_xml=None):
    import zipfile
    import io
    from xml.sax.saxutils import escape
    body = "".join((f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{escape(p[2:])}</w:t></w:r></w:p>' if p.startswith("# ")
                    else f"<w:p><w:r><w:t>{escape(p)}</w:t></w:r></w:p>") for p in paras)
    xml = document_xml or ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                           + body + "</w:body></w:document>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        if document_xml != "":
            z.writestr("word/document.xml", xml)
        for k, v in (extra_files or {}).items():
            z.writestr(k, v)
    return buf.getvalue()


def make_pdf(pages):
    """Hand-built text PDF (no libraries needed). pages = list of line lists."""
    esc = lambda l: l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    n = len(pages)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", ("<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join(f"{4 + 2 * i} 0 R" for i in range(n)), n)).encode(),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, lines in enumerate(pages):
        c = ("BT /F1 10 Tf 12 TL 40 760 Td " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET").encode("latin-1", "replace")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(c) + c + b"\nendstream")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()


def sample_pdf_pages():
    import textwrap
    lines = []
    for l in SAMPLE.split("\n"):
        lines += textwrap.wrap(l, 95) or [""]
    per = -(-len(lines) // 4)
    return [["CONFIDENTIAL HEADER ACME"] + lines[i * per:(i + 1) * per] + [f"Page {i + 1}"] for i in range(4)]


def call(method, path, body=None, raw=None):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(BASE + path, data=data, method=method, headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status, json.loads(r.read() or b"null"), r.headers
    except urllib.error.HTTPError as e:
        body = e.read()
        try:
            return e.code, json.loads(body), e.headers
        except ValueError:
            return e.code, body, e.headers


def setUpModule():
    global PROC, BASE
    tmp = tempfile.mkdtemp()
    db = os.path.join(tmp, "test.db")
    subprocess.run([sys.executable, os.path.join(HERE, "import_excel.py"), XLSX, db], check=True, capture_output=True)
    s = socket.socket(); s.bind(("127.0.0.1", 0)); port = s.getsockname()[1]; s.close()
    start_stub()
    env = dict(os.environ, ECOS_DB=db, PORT=str(port), ECOS_TODAY="2026-10-22", ECOS_QUIET="1", ANTHROPIC_API_KEY="sk-test-key",
               ANTHROPIC_BASE_URL=STUB_URL, ECOS_MAX_UPLOAD_MB="1", ECOS_RATE_LIMIT="1000", ECOS_SUMMARY_CONCURRENCY="8")
    PROC = subprocess.Popen([sys.executable, os.path.join(HERE, "server.py")], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    BASE = f"http://127.0.0.1:{port}"
    for _ in range(50):
        try:
            global BASELINE
            BASELINE = call("GET", "/api/health")[1]["counts"]
            global BASE_VIOL
            BASE_VIOL = sum(1 for c in call("GET", "/api/constraints?size=1000")[1]["data"] if c["Status"] == "Violated")
            return
        except Exception:
            time.sleep(.1)
    raise RuntimeError("server did not start")


def tearDownModule():
    PROC.terminate()


class Api(unittest.TestCase):
    def test_health_counts_match_workbook(self):
        self.assertEqual(call("GET", "/api/health")[0], 200)
        self.assertEqual(BASELINE, dict(employees=150, departments=12, projects=50, tasks=800, assignments=1200, constraints=400,
                                           communications=700, meetings=150, resources=200, dependencies=300))

    def test_bootstrap(self):
        s, d, _ = call("GET", "/api/bootstrap")
        self.assertEqual((s, len(d["Employees"]), len(d["Task_Dependencies"]), d["today"]), (200, 150, 300, "2026-10-22"))

    def test_list_filter_search_sort_paginate(self):
        s, d, _ = call("GET", "/api/tasks?Status=Pending&Priority=High,Critical&sort=Estimated_Hours&order=desc&size=5&page=2")
        self.assertEqual(s, 200)
        self.assertEqual(len(d["data"]), 5)
        self.assertTrue(all(t["Status"] == "Pending" and t["Priority"] in ("High", "Critical") for t in d["data"]))
        self.assertGreaterEqual(d["total"], 10)
        _, d2, _ = call("GET", "/api/employees?q=Aarav")
        self.assertTrue(d2["total"] > 0 and all("Aarav" in e["Employee_Name"] or "Aarav" in json.dumps(e) for e in d2["data"]))
        self.assertEqual(call("GET", "/api/tasks?sort=Bogus")[0], 400)
        self.assertEqual(call("GET", "/api/tasks?page=abc")[0], 400)

    def test_get_one_and_404(self):
        self.assertEqual(call("GET", "/api/employees/E001")[1]["Employee_Name"], "Aarav Sharma")
        s, d, _ = call("GET", "/api/employees/E999")
        self.assertEqual((s, d["error"]["code"]), (404, "not_found"))
        self.assertEqual(call("GET", "/api/nope")[0], 404)

    def test_employee_crud_lifecycle(self):
        good = dict(Employee_Name="Test Person", Department="QA", Role="QA Engineer", Skills="Testing,  SQL", Experience_Years=3,
                    Availability_Hours_Per_Day=6, Current_Workload_Percent=20, Performance_Score=80, Team_ID="T01", Work_Mode="Remote")
        s, d, _ = call("POST", "/api/employees", good)
        self.assertEqual(s, 201)
        eid = d["Employee_ID"]
        self.assertEqual((eid, d["Skills"]), ("E151", "Testing, SQL"))
        s, d, _ = call("PUT", f"/api/employees/{eid}", {"Current_Workload_Percent": 35})
        self.assertEqual((s, d["Current_Workload_Percent"], d["Employee_Name"]), (200, 35, "Test Person"))
        self.assertEqual(call("DELETE", f"/api/employees/{eid}")[0], 200)
        self.assertEqual(call("GET", f"/api/employees/{eid}")[0], 404)

    def test_validation_errors(self):
        s, d, _ = call("POST", "/api/employees", {"Employee_Name": "X", "Current_Workload_Percent": 250, "Work_Mode": "Moon"})
        self.assertEqual(s, 422)
        det = d["error"]["details"]
        self.assertIn("Current_Workload_Percent", det)
        self.assertIn("Work_Mode", det)
        self.assertIn("Department", det)  # missing required field
        self.assertEqual(call("POST", "/api/tasks", dict(Project_ID="P999", Task_Name="t", Required_Skills="SQL", Priority="Low", Estimated_Hours=3,
                                                         Deadline="2026-13-45", Required_Employees=1, Status="Pending", Department="QA"))[1]["error"]["details"].keys() >= {"Project_ID", "Deadline"}, True)
        self.assertEqual(call("PUT", "/api/employees/E001", {"Bogus": 1})[0], 400)
        self.assertEqual(call("PUT", "/api/employees/E001", {"Employee_ID": "E777"})[0], 422)
        self.assertEqual(call("PUT", "/api/employees/E001", {})[0], 400)
        self.assertEqual(call("POST", "/api/employees", raw=b"{not json")[0], 400)
        self.assertEqual(call("POST", "/api/employees", [1, 2])[0], 400)

    def test_task_crud_and_dependency_rules(self):
        t = dict(Project_ID="P001", Task_Name="API test task", Required_Skills="SQL", Priority="Low", Estimated_Hours=4,
                 Deadline="2026-12-01", Required_Employees=1, Status="Pending", Department="QA")
        s, d, _ = call("POST", "/api/tasks", t)
        self.assertEqual((s, d["Task_ID"], d["Dependency_Task_ID"]), (201, "T0801", "None"))
        self.assertEqual(call("PUT", "/api/tasks/T0801", {"Dependency_Task_ID": "T0801"})[0], 422)
        self.assertEqual(call("PUT", "/api/tasks/T0801", {"Dependency_Task_ID": "T0001"})[0], 200)
        self.assertEqual(call("PUT", "/api/tasks/T0801", {"Dependency_Task_ID": "T9999"})[0], 422)
        self.assertEqual(call("PUT", "/api/tasks/T0801", {"Status": "Completed"})[1]["Status"], "Completed")
        self.assertEqual(call("DELETE", "/api/tasks/T0801")[0], 200)

    def test_delete_blocked_by_references(self):
        s, d, _ = call("DELETE", "/api/projects/P001")
        self.assertEqual((s, d["error"]["code"]), (409, "in_use"))
        self.assertEqual(call("GET", "/api/projects/P001")[0], 200)

    def test_duplicate_primary_key(self):
        s, d, _ = call("POST", "/api/departments", {"Department_ID": "D01", "Department_Name": "Dup", "Typical_Team_Size": 5, "Department_Type": "Support"})
        self.assertEqual((s, d["error"]["code"]), (409, "duplicate"))
        s, d, _ = call("POST", "/api/departments", {"Department_Name": "Legal", "Typical_Team_Size": 5, "Department_Type": "Support"})
        self.assertEqual((s, d["Department_ID"]), (201, "D13"))

    def test_meeting_and_message_validation(self):
        m = dict(Project_ID="P001", Participants="E001,E002", Topic="Sync", Discussion_Summary="s", Decisions="d", Action_Items="a", Meeting_Date="2026-11-01")
        s, d, _ = call("POST", "/api/meetings", m)
        self.assertEqual((s, d["Participants"], d["Meeting_ID"]), (201, "E001, E002", "MT151"))
        self.assertEqual(call("POST", "/api/meetings", {**m, "Participants": "E001,E999"})[0], 422)
        msg = dict(Sender_ID="E001", Receiver_ID="E002", Project_ID="P001", Communication_Type="Chat", Message_Text="hi",
                   Timestamp="2026-10-22 09:30", Priority="Low", Action_Required="No")
        self.assertEqual(call("POST", "/api/communications", msg)[0], 201)
        self.assertEqual(call("POST", "/api/communications", {**msg, "Timestamp": "yesterday"})[0], 422)

    def test_constraints_have_computed_status(self):
        _, d, _ = call("GET", "/api/constraints?size=1000")
        self.assertEqual(d["total"], 400)
        self.assertEqual(BASE_VIOL, 52)  # same result as the in-browser evaluation of the workbook
        self.assertTrue(all(c["Status"] in ("Satisfied", "Violated") for c in d["data"]))
        _, v, _ = call("GET", "/api/constraints/validate?task_id=" + d["data"][0]["Task_ID"])
        self.assertTrue(v["total"] >= 1)
        self.assertEqual(call("GET", "/api/constraints/validate?task_id=T9999")[0], 404)

    def test_optimize(self):
        s, d, _ = call("POST", "/api/optimize", {"task_id": "T0001", "top": 5})
        self.assertEqual(s, 200)
        c = d["candidates"]
        self.assertEqual((len(c), d["evaluated"]), (5, 150 - len(d["assigned"])))  # employees already on the task are excluded
        self.assertEqual([x["rank"] for x in c], [1, 2, 3, 4, 5])
        passing = [x["passes_hard_constraints"] for x in c]
        self.assertEqual(passing, sorted(passing, reverse=True))
        scores = [x["score"] for x in c if x["passes_hard_constraints"]]
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertEqual(d["best"]["employee_id"], c[0]["employee_id"])
        self.assertEqual(call("POST", "/api/optimize", {"task_id": "T9999"})[0], 404)
        self.assertEqual(call("POST", "/api/optimize", {"task_id": "T0001", "top": "x"})[0], 422)
        self.assertEqual(call("GET", "/api/optimize")[0], 405)

    def test_apply_assignment_flow(self):
        _, d, _ = call("POST", "/api/optimize", {"task_id": "T0001"})
        self.assertTrue(1 <= len(d["options"]) <= 3)
        emp = d["options"][0]["employee_id"]
        s, a, _ = call("POST", "/api/assignments/apply", {"task_id": "T0001", "employee_id": emp})
        self.assertEqual(s, 201)
        self.assertEqual((a["assignment"]["Assignment_Status"], a["assignment"]["Employee_ID"]), ("Assigned", emp))
        self.assertEqual(a["task"]["Status"], "In Progress")  # was Pending
        s, e, _ = call("POST", "/api/assignments/apply", {"task_id": "T0001", "employee_id": emp})
        self.assertEqual((s, e["error"]["code"]), (409, "duplicate_assignment"))
        self.assertEqual(call("POST", "/api/assignments/apply", {"task_id": "T0001"})[0], 422)
        self.assertEqual(call("POST", "/api/assignments/apply", {"task_id": "T0001", "employee_id": "E999"})[0], 404)

    def test_three_free_options_and_one_task_per_person(self):
        _, d, _ = call("POST", "/api/optimize", {"task_id": "T0002"})
        opts = d["options"]
        self.assertEqual(len(opts), 3)
        self.assertEqual(len({o["employee_id"] for o in opts}), 3)
        self.assertTrue(all(o["available"] and o["open_tasks"] == 0 and o["passes_hard_constraints"] and o["availability_hours"] >= 5 for o in opts))
        self.assertEqual([o["option"] for o in opts], [1, 2, 3])
        scores = [o["score"] for o in opts]
        self.assertEqual(scores, sorted(scores, reverse=True))
        pick = opts[1]["employee_id"]  # the user chooses option 2
        s, r, _ = call("POST", "/api/assignments/apply", {"task_id": "T0002", "employee_id": pick})
        self.assertEqual((s, len(r["assignments"])), (201, 1))
        # the chosen person is now busy: never offered again, and cannot take a second open task
        _, other, _ = call("GET", "/api/tasks?Status=Pending&size=40")
        for t_ in other["data"]:
            if t_["Task_ID"] == "T0002":
                continue
            _, o2, _ = call("POST", "/api/optimize", {"task_id": t_["Task_ID"], "options": 10})
            self.assertNotIn(pick, [x["employee_id"] for x in o2["options"]])
        tid2 = next(x["Task_ID"] for x in other["data"] if x["Task_ID"] != "T0002")
        s, e, _ = call("POST", "/api/assignments/apply", {"task_id": tid2, "employee_id": pick})
        self.assertIn((s, e["error"]["code"]), [(409, "employee_busy"), (409, "constraint_violation")])

    def test_apply_blocks_hard_constraint_violation(self):
        _, cons, _ = call("GET", "/api/constraints?size=1000&Constraint_Level=Hard&Constraint_Type=Minimum%20Experience")
        found = False
        for k in cons["data"]:
            _, d, _ = call("POST", "/api/optimize", {"task_id": k["Task_ID"], "top": 150})
            bad = [x for x in d["candidates"] if not x["passes_hard_constraints"]]
            if bad:
                s, e, _ = call("POST", "/api/assignments/apply", {"task_id": k["Task_ID"], "employee_id": bad[0]["employee_id"]})
                self.assertEqual((s, e["error"]["code"]), (409, "constraint_violation"))
                self.assertTrue(e["error"]["details"]["violations"])
                s, ok, _ = call("POST", "/api/assignments/apply", {"task_id": k["Task_ID"], "employee_id": bad[0]["employee_id"], "force": True})
                self.assertEqual(s, 201)
                self.assertTrue(ok["warnings"])
                found = True
                break
        self.assertTrue(found)

    def test_staffing_a_multi_person_task_one_at_a_time(self):
        _, tl, _ = call("GET", "/api/tasks?Required_Employees=3&Status=Pending&size=300")
        tid = next(x["Task_ID"] for x in tl["data"] if call("GET", f"/api/assignments?Task_ID={x['Task_ID']}&Assignment_Status=Assigned")[1]["total"] == 0)
        chosen = []
        for need in (3, 2, 1):
            _, o, _ = call("POST", "/api/optimize", {"task_id": tid})
            self.assertEqual((o["required"], o["remaining"], len(o["options"])), (3, need, 3))
            self.assertFalse(set(chosen) & {x["employee_id"] for x in o["options"]})
            pick = o["options"][-1]["employee_id"]  # last option, to prove any of the three can be chosen
            s, r, _ = call("POST", "/api/assignments/apply", {"task_id": tid, "employee_id": pick})
            self.assertEqual((s, len(r["assignments"]), r["still_needed"]), (201, 1, need - 1))
            chosen.append(pick)
        _, o, _ = call("POST", "/api/optimize", {"task_id": tid})
        self.assertEqual((o["remaining"], o["options"], len(o["assigned"])), (0, [], 3))
        s, e, _ = call("POST", "/api/assignments/apply", {"task_id": tid, "auto": True})
        self.assertEqual((s, e["error"]["code"]), (409, "no_candidates"))
        extra = next(x["employee_id"] for x in o["candidates"] if x["available"])
        s, e, _ = call("POST", "/api/assignments/apply", {"task_id": tid, "employee_ids": [extra]})
        self.assertEqual((s, e["error"]["code"]), (409, "task_full"))

    def test_auto_assigns_one_best_free_person(self):
        _, tl, _ = call("GET", "/api/tasks?Required_Employees=2&Status=Blocked&size=300")
        tid = next(x["Task_ID"] for x in tl["data"] if call("GET", f"/api/assignments?Task_ID={x['Task_ID']}&Assignment_Status=Assigned")[1]["total"] == 0)
        _, o, _ = call("POST", "/api/optimize", {"task_id": tid})
        s, r, _ = call("POST", "/api/assignments/apply", {"task_id": tid, "auto": True})
        self.assertEqual((s, len(r["assignments"]), r["assignment"]["Employee_ID"]), (201, 1, o["options"][0]["employee_id"]))
        self.assertEqual(r["task"]["Status"], "Blocked")  # only Pending tasks move to In Progress

    def test_update_suggested_assignment(self):
        _, d, _ = call("GET", "/api/assignments?Assignment_Status=Suggested&size=1")
        a = d["data"][0]
        s, r, _ = call("PUT", f'/api/assignments/{a["Assignment_ID"]}', {"Assignment_Status": "Assigned"})
        self.assertEqual((s, r["Assignment_Status"]), (200, "Assigned"))
        self.assertEqual(call("PUT", f'/api/assignments/{a["Assignment_ID"]}', {"Skill_Match_Score": 5})[0], 422)

    def test_reports(self):
        s, d, _ = call("GET", "/api/dashboard")
        self.assertEqual(s, 200)
        self.assertEqual(d["total_employees"], 150)
        self.assertEqual(len(d["insights"]), 4)
        _, a, _ = call("GET", "/api/analytics")
        self.assertTrue(0 <= a["collaboration_optimization_score"] <= 100)
        self.assertIn("Software Development", a["department_productivity"])
        _, r, _ = call("GET", "/api/recommendations")
        self.assertEqual(set(r), {"task", "workforce", "deadline_risk", "resource"})
        self.assertTrue(r["resource"] and r["workforce"])

    def test_employee_report(self):
        s, r, _ = call("GET", "/api/employees/E001/report")
        self.assertEqual(s, 200)
        self.assertEqual(set(r), {"generated_for", "employee", "summary", "assignments_by_status", "current_tasks", "completed_tasks", "suggested_tasks",
                                  "projects", "comparison", "communications", "meetings", "insights"})
        _, emp, _ = call("GET", "/api/employees/E001")
        self.assertEqual((r["employee"]["name"], r["employee"]["department"], r["employee"]["performance_score"]),
                         (emp["Employee_Name"], emp["Department"], emp["Performance_Score"]))
        self.assertEqual(r["generated_for"], "2026-10-22")
        # counts agree with the raw tables
        _, asg, _ = call("GET", "/api/assignments?Employee_ID=E001&size=1000")
        self.assertEqual(r["summary"]["total_assignments"], asg["total"])
        self.assertEqual(sum(r["assignments_by_status"].values()), asg["total"])
        _, msgs, _ = call("GET", "/api/communications?Sender_ID=E001&size=1000")
        self.assertEqual(r["summary"]["messages_sent"], msgs["total"])
        _, got, _ = call("GET", "/api/communications?Receiver_ID=E001&size=1000")
        self.assertEqual(r["summary"]["messages_received"], got["total"])
        _, mt, _ = call("GET", "/api/meetings?size=1000")
        self.assertEqual(r["summary"]["meetings_attended"], sum(1 for m in mt["data"] if "E001" in [x.strip() for x in m["Participants"].split(",")]))
        self.assertTrue(r["insights"])
        self.assertEqual(r["comparison"]["organization_size"], 150)

    def test_employee_report_open_tasks_and_free_flag_match_the_optimizer(self):
        _, asg, _ = call("GET", "/api/assignments?Assignment_Status=Assigned&size=1000")
        _, tks, _ = call("GET", "/api/tasks?size=1000")
        open_ids = {t["Task_ID"] for t in tks["data"] if t["Status"] != "Completed"}
        who = next(a["Employee_ID"] for a in asg["data"] if a["Task_ID"] in open_ids)
        mine = {a["Task_ID"] for a in asg["data"] if a["Employee_ID"] == who and a["Task_ID"] in open_ids}
        s, r, _ = call("GET", f"/api/employees/{who}/report")
        self.assertEqual(s, 200)
        self.assertEqual({t["task_id"] for t in r["current_tasks"]}, mine)
        self.assertEqual(r["summary"]["open_tasks"], len(mine))
        self.assertFalse(r["employee"]["free"])
        for t in r["current_tasks"]:
            self.assertEqual(t["days_left"] < 0, t["task_id"] in {x["task_id"] for x in r["current_tasks"] if x["deadline"] < "2026-10-22"})
        # an employee with no open task and >=5 h/day is reported as free
        busy = {a["Employee_ID"] for a in asg["data"] if a["Task_ID"] in open_ids}
        _, emps, _ = call("GET", "/api/employees?size=1000")
        idle = next(e["Employee_ID"] for e in emps["data"] if e["Employee_ID"] not in busy and e["Availability_Hours_Per_Day"] >= 5)
        self.assertTrue(call("GET", f"/api/employees/{idle}/report")[1]["employee"]["free"])

    def test_employee_report_errors_and_routing(self):
        self.assertEqual(call("GET", "/api/employees/E999/report")[0], 404)
        self.assertEqual(call("POST", "/api/employees/E001/report", {})[0], 405)
        self.assertEqual(call("GET", "/api/employees/E001")[0], 200)  # the plain record route still works
        self.assertEqual(call("GET", "/api/employees/E001/report/extra")[0], 405)  # same answer as any other unknown sub-path
        s, r, _ = call("GET", "/api/employees/%20E001%20/report")
        self.assertEqual(s, 404)

    def test_every_employee_has_a_report(self):
        _, emps, _ = call("GET", "/api/employees?size=1000")
        for e in emps["data"]:
            s, r, _ = call("GET", f'/api/employees/{e["Employee_ID"]}/report')
            self.assertEqual(s, 200, e["Employee_ID"])
            self.assertEqual(r["employee"]["employee_id"], e["Employee_ID"])

    def test_meeting_summary(self):
        s, d, _ = call("POST", "/api/meetings/MT001/summary", {})
        self.assertEqual(s, 200)
        self.assertIn("Requirements Discussion", d["summary"])
        self.assertEqual(call("POST", "/api/meetings/MT999/summary", {})[0], 404)

    def test_cors_static_and_traversal(self):
        req = urllib.request.Request(BASE + "/api/tasks", method="OPTIONS")
        with urllib.request.urlopen(req) as r:
            self.assertEqual((r.status, r.headers["Access-Control-Allow-Origin"]), (204, "*"))
        for p in ("/..%2f..%2fetc%2fpasswd", "/../backend/server.py", "/%2e%2e/backend/ecos.db"):
            code = call("GET", p)[0]
            self.assertEqual(code, 404, p)


REPORT_KEYS = {"overview", "key_points", "findings", "statistics", "conclusions", "recommendations", "meta"}


class Reports(unittest.TestCase):
    def check_good(self, r, method="local"):
        self.assertEqual(set(r), REPORT_KEYS)
        self.assertTrue(r["overview"].strip())
        self.assertEqual(r["meta"]["method"], method)
        self.assertTrue(any("recommend" in x.lower() or "should" in x.lower() or "must" in x.lower() for x in r["recommendations"]))
        self.assertTrue(any("62%" in x for x in r["statistics"]))
        self.assertTrue(any("conclusion" in x.lower() or "driver" in x.lower() for x in r["conclusions"]))
        self.assertGreaterEqual(len(r["key_points"]), 3)
        self.assertNotIn("maria.silva@example.com", json.dumps(r))  # contact details never appear in summaries
        self.assertNotIn("415-555-0199", json.dumps(r))

    def test_config_hides_secrets(self):
        s, d, _ = call("GET", "/api/reports/config")
        self.assertEqual((s, d["ai_available"], d["max_mb"], d["formats"]), (200, True, 1.0, [".pdf", ".docx", ".txt"]))
        self.assertNotIn("sk-test-key", json.dumps(d))
        self.assertEqual(call("GET", "/api/reports/summarize")[0], 405)

    def test_txt_local(self):
        s, r = upload("q3 report.txt", SAMPLE.encode())
        self.assertEqual(s, 200)
        self.check_good(r)
        self.assertEqual((r["meta"]["filename"], r["meta"]["type"]), ("q3 report.txt", "txt"))
        self.assertTrue(r["meta"]["words"] > 250)

    def test_txt_encodings(self):
        for enc in ("utf-8-sig", "utf-16", "cp1252"):
            s, r = upload("r.txt", SAMPLE.replace("organisation", "organisation \u2013 note").encode(enc))
            self.assertEqual(s, 200, enc)
            self.check_good(r)

    def test_docx_local(self):
        paras = [("# " + l if l and l[:1].isupper() and len(l) < 40 and not l.endswith(".") else l) for l in SAMPLE.split("\n") if l.strip()]
        s, r = upload("report.docx", make_docx(paras))
        self.assertEqual(s, 200)
        self.check_good(r)
        self.assertEqual(r["meta"]["type"], "docx")

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed (only used to create a test file)")
    def test_docx_made_by_word_library(self):
        import docx
        import io
        d = docx.Document()
        for l in SAMPLE.split("\n"):
            if l.strip():
                (d.add_heading(l, 1) if len(l) < 40 and not l.endswith(".") else d.add_paragraph(l))
        tbl = d.add_table(rows=1, cols=2)
        tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "Metric", "Value"
        buf = io.BytesIO()
        d.save(buf)
        s, r = upload("word.docx", buf.getvalue())
        self.assertEqual(s, 200)
        self.check_good(r)

    def test_pdf_local_with_headers_and_footers(self):
        pdf = make_pdf(sample_pdf_pages())
        s, r = upload("report.pdf", pdf)
        self.assertEqual(s, 200)
        self.check_good(r)
        self.assertEqual((r["meta"]["pages"], r["meta"]["type"]), (4, "pdf"))
        self.assertNotIn("CONFIDENTIAL HEADER", json.dumps(r))  # repeated header was stripped

    @unittest.skipUnless(importlib.util.find_spec("reportlab"), "reportlab not installed (only used to create a test file)")
    def test_pdf_made_by_reportlab(self):
        import io
        import textwrap
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        y = 800
        for l in SAMPLE.split("\n"):
            for w in textwrap.wrap(l, 90) or [""]:
                c.drawString(40, y, w)
                y -= 14
                if y < 60:
                    c.showPage()
                    y = 800
        c.save()
        s, r = upload("rl.pdf", buf.getvalue())
        self.assertEqual(s, 200)
        self.check_good(r)

    def test_pdf_problems(self):
        self.assertEqual(upload("a.pdf", b"%PDF-1.4\nthis is not really a pdf")[1]["error"]["code"], "unreadable")
        empty = upload("scan.pdf", make_pdf([["" ]]))
        self.assertEqual((empty[0], empty[1]["error"]["code"]), (422, "no_text"))
        try:
            import io
            from pypdf import PdfReader, PdfWriter
            w = PdfWriter()
            for p in PdfReader(io.BytesIO(make_pdf(sample_pdf_pages()))).pages:
                w.add_page(p)
            w.encrypt("secret", algorithm="RC4-128")
            buf = io.BytesIO()
            w.write(buf)
        except Exception:
            self.skipTest("could not create an encrypted PDF here")
        s, r = upload("locked.pdf", buf.getvalue())
        self.assertEqual((s, r["error"]["code"]), (422, "password_protected"))

    def test_upload_validation(self):
        cases = [("x.exe", b"MZ\x90\x00data", 415, "unsupported_type"), ("x.doc", b"\xd0\xcf\x11\xe0", 415, "unsupported_type"), ("noext", b"hello world", 415, "unsupported_type"),
                 ("x.pdf", b"just plain text " * 20, 415, "type_mismatch"), ("x.txt", b"%PDF-1.4 hello", 415, "type_mismatch"),
                 ("x.txt", b"abc\x00\x01\x02binary" * 50, 415, "type_mismatch"), ("x.docx", b"plain text not zip", 415, "type_mismatch"),
                 ("x.txt", b"", 400, "empty_file"), ("x.txt", b"too short to summarize", 422, "no_text"),
                 ("x.docx", b"PK\x03\x04junk-not-a-zip", 422, "unreadable"), ("x.docx", make_docx(["hi"], document_xml=""), 422, "unreadable")]
        for name, data, status, code in cases:
            s, r = upload(name, data)
            self.assertEqual((s, r["error"]["code"]), (status, code), name)
        bomb = make_docx([], document_xml='<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        self.assertEqual(upload("e.docx", bomb)[1]["error"]["code"], "unsafe_file")
        self.assertEqual(upload("x.txt", SAMPLE.encode(), mode="bogus")[1]["error"]["code"], "bad_mode")

    def test_size_limit_and_server_stays_healthy(self):
        s, r = upload("big.txt", b"word " * 250_000)  # 1.25 MB > the 1 MB test limit
        self.assertEqual((s, r["error"]["code"]), (413, "too_large"))
        self.assertEqual(call("GET", "/api/health")[0], 200)
        s, r = upload("ok.txt", SAMPLE.encode())
        self.assertEqual(s, 200)

    def test_filename_is_sanitised(self):
        s, r = upload("../../etc/pass<wd>.txt", SAMPLE.encode(), extra_headers={"X-File-Name": urllib.parse.quote("..\\..\\etc/pass<wd>.txt")})
        self.assertEqual((s, r["meta"]["filename"]), (200, "passwd.txt"))

    def test_nothing_is_stored_on_disk(self):
        before = set(os.listdir(HERE))
        upload("keep.txt", SAMPLE.encode())
        self.assertEqual(set(os.listdir(HERE)), before)
        self.assertNotIn("Report", json.dumps(call("GET", "/api/bootstrap")[1])[:0] + "")

    def test_ai_summary_masks_private_data_and_verifies_figures(self):
        STUB_LOG.clear()
        s, r = upload("ai.txt", SAMPLE.encode(), mode="ai")
        self.assertEqual(s, 200)
        self.assertEqual((r["meta"]["method"], r["meta"]["model"]), ("ai", "claude-sonnet-5-5"))
        self.assertIn("Resource contention", " ".join(r["conclusions"]))
        self.assertTrue(r["meta"]["masked_items"] >= 2)
        self.assertNotIn("999%", json.dumps(r))  # invented figure is dropped
        self.assertTrue(any("removed" in n for n in r["meta"]["notes"]))
        sent = STUB_LOG[-1]
        self.assertEqual({k.lower(): v for k, v in sent["headers"].items()}.get("x-api-key"), "sk-test-key")
        self.assertIn("untrusted", sent["body"]["system"])
        self.assertIn("[EMAIL]", sent["doc"])
        self.assertIn("[PHONE]", sent["doc"])
        self.assertNotIn("maria.silva@example.com", sent["doc"])
        self.assertNotIn("415-555-0199", sent["doc"])

    def test_ai_failures(self):
        for marker, code in (("TRIGGER_FAIL", "ai_unavailable"), ("TRIGGER_BADJSON", "ai_bad_output"), ("TRIGGER_AUTH", "ai_unavailable")):
            doc = (SAMPLE + " " + marker).encode()
            s, r = upload("f.txt", doc, mode="ai")
            self.assertEqual((s, r["error"]["code"]), (502, code), marker)
            s, r = upload("f.txt", doc, mode="auto")  # automatic mode falls back to the built-in summarizer
            self.assertEqual((s, r["meta"]["method"]), (200, "local"))
            self.assertTrue(any("AI summary unavailable" in n for n in r["meta"]["notes"]))
        self.assertNotIn("sk-test-key", json.dumps(r))

    def test_long_reports_are_condensed_for_ai(self):
        import sys
        sys.path.insert(0, HERE)
        import summarizer
        big = SAMPLE + "\nAppendix A\n" + "\n".join(f"Appendix note {i}: batch {i * 7} of the sensor array was calibrated by technician {chr(65 + i % 26)}{i} during shift {i % 3}." for i in range(1, 400))
        small = summarizer.condense(big, 20000)
        self.assertLessEqual(len(small), 20000)
        self.assertIn("recommend", small)
        self.assertEqual(summarizer.condense("short text", 100), "short text")

    def test_rate_limit_and_missing_key_unit(self):
        import sys
        sys.path.insert(0, HERE)
        import summarizer
        old = dict(os.environ)
        try:
            os.environ["ECOS_RATE_LIMIT"] = "3"
            for _ in range(3):
                summarizer.check_rate("9.9.9.9")
            with self.assertRaises(summarizer.SummaryError) as cm:
                summarizer.check_rate("9.9.9.9")
            self.assertEqual(cm.exception.status, 429)
            summarizer.check_rate("8.8.8.8")  # other clients are unaffected
            os.environ["ECOS_RATE_LIMIT"] = "1000"
            os.environ.pop("ANTHROPIC_API_KEY", None)
            with self.assertRaises(summarizer.SummaryError) as cm:
                summarizer.summarize_upload(SAMPLE.encode(), "a.txt", "ai", "1.1.1.1")
            self.assertEqual(cm.exception.code, "ai_not_configured")
            self.assertEqual(summarizer.summarize_upload(SAMPLE.encode(), "a.txt", "auto", "1.1.1.1")["meta"]["method"], "local")
        finally:
            os.environ.clear()
            os.environ.update(old)

    def test_redaction(self):
        import sys
        sys.path.insert(0, HERE)
        import summarizer
        pre = "23456789012"
        ok = pre + next(str(d) for d in range(10) if summarizer.verhoeff_valid(pre + str(d)))
        aad = f"{ok[:4]} {ok[4:8]} {ok[8:]}"
        txt, n = summarizer.redact(f"Mail a.b@corp.io, call +1 415-555-0199 or (415) 555-0100. SSN 123-45-6789. Card 4111 1111 1111 1111. Order 4111 1111 1111 1112. Aadhaar {aad}. Sales hit 2,500,000 in 2023 2024 2025.")
        for bad in ("a.b@corp.io", "415-555-0199", "123-45-6789", "4111 1111 1111 1111", aad):
            self.assertNotIn(bad, txt)
        self.assertIn("4111 1111 1111 1112", txt)  # fails the card checksum, so it is not a card
        self.assertIn("2,500,000 in 2023 2024 2025", txt)  # ordinary figures are kept
        self.assertEqual(n, 6)


if __name__ == "__main__":
    unittest.main(verbosity=2)
